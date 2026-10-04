from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook

from reefvision.storage.results_packages import (
    build_results_package,
    safe_component_limit,
    safe_package_name,
    verify_results_package,
)


def _run(*, tuning: bool = False) -> dict[str, object]:
    epoch_rows = [
        {
            "epoch": epoch,
            "train/box_loss": 2.0 / (epoch + 1),
            "metrics/precision(B)": 0.5 + epoch / 1000,
            "metrics/mAP50-95(B)": 0.2 + epoch / 2000,
            "lr/pg0": 0.01 / (epoch + 1),
        }
        for epoch in range(205)
    ]
    trials = [
        {
            "trial_number": index + 1,
            "status": "failed" if index == 17 else "completed",
            "fitness": None if index == 17 else 0.1 + ((index * 13) % 97) / 1000,
            "metrics": {"map50_95": 0.2 + index / 2000},
            "hyperparameters": {f"parameter_{item}": index / (item + 2) for item in range(21)},
            "duration_seconds": 3.5 + index,
            "epoch_history": [
                {"epoch": epoch, "train_loss": 1 / (epoch + 1), "map50_95": 0.2 + epoch / 100}
                for epoch in range(3)
            ]
            if index < 2
            else [],
            "artifacts": [],
        }
        for index in range(208)
    ]
    return {
        "id": "12345678-1234-1234-1234-123456789abc",
        "run_name": "unsafe/name:=formula",
        "status": "completed",
        "stage": "completed",
        "task": "detect",
        "provider_id": "ultralytics",
        "architecture": "yolo11n.yaml",
        "dataset_version_id": "dataset-id",
        "dataset_manifest_sha256": "a" * 64,
        "split_manifest_sha256": "b" * 64,
        "initial_checkpoint_source": "scratch",
        "initial_checkpoint_model_id": None,
        "resolved_model_asset_id": "builtin_architecture:yolo11n.yaml",
        "resolved_model_checksum": None,
        "device": "cpu",
        "precision": "fp32",
        "random_seed": 42,
        "training_config": {"epochs": 205, "name": "=unsafe"},
        "tuning_config": {"mode": "random" if tuning else "disabled"},
        "evaluation_config": {"split": "validation"},
        "tiling_config": {"mode": "none"},
        "resource_limits": {},
        "preflight_report": {"resource_snapshot": {"cpu": "fixture"}, "api_key": "never-export"},
        "draft_snapshot": {},
        "metrics": {"map50_95": 0.40123456789},
        "metrics_history": {
            "epoch": list(range(205)),
            "metrics/mAP50-95(B)": [row["metrics/mAP50-95(B)"] for row in epoch_rows],
        },
        "native_results_rows": epoch_rows,
        "tuning_results": {
            "fitness_definition": "0.1 x mAP50 + 0.9 x mAP50-95",
            "search_space": {"lr0": {"enabled": True, "min": 1e-5, "max": 0.1}},
            "max_trials": 250,
            "epochs_per_trial": 3,
            "patience": 50,
            "trials": trials,
            "best_trial_number": 97,
            "best_hyperparameters": trials[96]["hyperparameters"],
        }
        if tuning
        else None,
        "result_checkpoint_path": None,
        "warnings": [],
        "created_at": "2026-08-26T01:00:00+00:00",
        "started_at": "2026-08-26T01:01:00+00:00",
        "completed_at": "2026-08-26T02:01:00+00:00",
    }


def _build(tmp_path: Path, *, tuning: bool = False) -> Path:
    artifact = tmp_path / "native" / "results.png"
    artifact.parent.mkdir()
    artifact.write_bytes(b"native-provider-plot")
    result = build_results_package(
        export_root=tmp_path / "exports",
        run=_run(tuning=tuning),
        project={"id": "project-id", "name": "Fixture Project", "slug": "fixture"},
        dataset={
            "id": "dataset-id",
            "name": "Dataset 1",
            "class_map": [{"class_index": 0, "name": "=unsafe class"}],
        },
        dataset_items=[
            {
                "image_id": "image-id",
                "original_filename": "frame.jpg",
                "split": "train",
                "group_key": "fragment-1",
            }
        ],
        source_artifacts=[({"name": "results.png", "category": "plot"}, artifact)],
        options={
            "include_excel": True,
            "include_machine_data": True,
            "include_native_artifacts": True,
            "include_logs": True,
            "checkpoint_mode": "none",
        },
    )
    return tmp_path / "exports" / str(result["package_name"])


def test_training_package_preserves_epoch_values_and_verifies(tmp_path: Path) -> None:
    package = _build(tmp_path)
    workbook_path = next(package.glob("*.xlsx"))
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    assert "Epoch Metrics" in workbook.sheetnames
    sheet = workbook["Epoch Metrics"]
    assert sheet.max_row == 206
    assert isinstance(sheet.cell(2, 2).value, (int, float))
    assert sheet.cell(2, 2).value == 2.0
    workbook.close()
    manifest = json.loads((package / "manifest.json").read_text(encoding="utf-8"))
    assert manifest["reefvision_results_package_schema"] == "1.0"
    assert all(not Path(item["relative_path"]).is_absolute() for item in manifest["files"])
    assert verify_results_package(package)["valid"] is True
    provenance = (package / "data" / "provenance.json").read_text(encoding="utf-8")
    assert "never-export" not in provenance


def test_tuning_package_retains_failed_and_multi_epoch_trials(tmp_path: Path) -> None:
    package = _build(tmp_path, tuning=True)
    workbook = load_workbook(next(package.glob("*.xlsx")), read_only=True, data_only=True)
    assert workbook["All Trials"].max_row == 209
    assert workbook["Trial Epoch Metrics"].max_row == 7
    statuses = [row[2].value for row in workbook["All Trials"].iter_rows(min_row=2)]
    assert "failed" in statuses
    workbook.close()
    assert (package / "data" / "parameter_fitness.csv").is_file()
    assert (package / "plots" / "reefvision" / "fitness_vs_trial.svg").is_file()


def test_verification_detects_tampering_and_names_are_windows_safe(tmp_path: Path) -> None:
    package = _build(tmp_path)
    target = package / "data" / "epoch_metrics.csv"
    target.write_text("tampered", encoding="utf-8")
    result = verify_results_package(package)
    assert result["valid"] is False
    assert any("mismatch" in error.casefold() for error in result["errors"])
    assert safe_package_name('bad<>:"/\\|?* name.').endswith("name")


def test_path_component_limit_reserves_nested_package_headroom() -> None:
    parent = Path("C:/") / ("nested-" * 15)

    limit = safe_component_limit(
        parent,
        requested=90,
        reserve=72,
    )

    assert limit >= 8
    assert limit < 90

    candidate_length = len(str(parent)) + 1 + limit + 72

    assert candidate_length <= 230
