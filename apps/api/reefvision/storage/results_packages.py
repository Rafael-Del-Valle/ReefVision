"""Atomic, safe scientific Results Package construction and verification."""

from __future__ import annotations

import csv
import hashlib
import json
import mimetypes
import re
import shutil
import statistics
import tempfile
from collections.abc import Iterable, Mapping, Sequence
from datetime import UTC, datetime
from pathlib import Path
from typing import Any
from xml.sax.saxutils import escape

import yaml
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

PACKAGE_SCHEMA_VERSION = "1.0"
CHUNK_SIZE = 1024 * 1024

# Keep generated package paths comfortably below the classic Windows
# MAX_PATH boundary. The budget is intentionally conservative because
# packages contain nested data/, plots/, artifacts/, and workbook paths.
RESULTS_PACKAGE_PATH_BUDGET = 230
RESULTS_PACKAGE_MIN_COMPONENT = 8
SECRET_KEYS = re.compile(r"(?:api[_-]?key|token|password|secret|credential)", re.IGNORECASE)


def safe_package_name(value: str, *, limit: int = 90) -> str:
    cleaned = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", value).strip(" .")
    cleaned = re.sub(r"\s+", "_", cleaned)
    return cleaned[:limit].rstrip(" ._") or "run"


def safe_component_limit(
    parent: Path,
    *,
    requested: int,
    reserve: int = 0,
) -> int:
    """Limit one generated path component while reserving nested-path space."""

    available = RESULTS_PACKAGE_PATH_BUDGET - len(str(parent)) - 1 - max(0, reserve)

    return max(
        RESULTS_PACKAGE_MIN_COMPONENT,
        min(requested, available),
    )


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        while chunk := stream.read(CHUNK_SIZE):
            digest.update(chunk)
    return digest.hexdigest()


def build_results_package(
    *,
    export_root: Path,
    run: Mapping[str, Any],
    project: Mapping[str, Any],
    dataset: Mapping[str, Any] | None,
    dataset_items: Sequence[Mapping[str, Any]],
    source_artifacts: Sequence[tuple[Mapping[str, Any], Path]],
    options: Mapping[str, Any],
) -> dict[str, Any]:
    tuning = isinstance(run.get("tuning_results"), Mapping)
    package_type = "tuning" if tuning else "training"
    timestamp = datetime.now(UTC)
    export_root.mkdir(parents=True, exist_ok=True)

    stem_limit = safe_component_limit(
        export_root,
        requested=90,
        reserve=72,
    )

    stem = safe_package_name(
        (f"ReefVision_{package_type.title()}_Run_{run['run_name']}_{timestamp:%Y%m%dT%H%M%SZ}"),
        limit=stem_limit,
    )

    final_dir = export_root / stem

    if final_dir.exists() or final_dir.with_suffix(".zip").exists():
        stem = safe_package_name(
            (f"rv_{str(run['id'])[:8]}_{timestamp.microsecond:06d}"),
            limit=stem_limit,
        )
        final_dir = export_root / stem

    warnings = list(run.get("warnings", []))

    # The staging directory deliberately does not repeat the final,
    # descriptive package stem. On Windows, export_root + temporary
    # directory + final stem + nested package content can exceed
    # MAX_PATH even when the final package location itself is valid.
    with tempfile.TemporaryDirectory(
        prefix="rv-",
        dir=export_root,
    ) as temporary:
        package = Path(temporary) / "p"
        package.mkdir()
        data_dir = package / "data"
        data_dir.mkdir()
        copied = _write_common_data(package, run, project, dataset, dataset_items, warnings)
        if options.get("include_machine_data", True):
            if tuning:
                copied.extend(_write_tuning_data(data_dir, run, warnings))
            else:
                copied.extend(_write_training_data(data_dir, run, warnings))
        copied.extend(_write_standard_plots(package, run))
        copied.extend(_copy_artifacts(package, source_artifacts, options, warnings))
        if options.get("include_machine_data", True):
            inventory = data_dir / "artifacts.csv"
            _write_csv(
                inventory,
                [
                    {
                        "artifact_name": item.name,
                        "category": category,
                        "relative_package_path": item.relative_to(package).as_posix(),
                        "source_type": source_type,
                        "file_size_bytes": item.stat().st_size,
                        "sha256": sha256_file(item),
                    }
                    for item, category, source_type in copied
                    if item.is_file()
                ],
            )
            copied.append((inventory, "data", "reefvision_generated"))
        if options.get("include_excel", True):
            workbook_stem_limit = safe_component_limit(
                package,
                requested=90,
                reserve=len(".xlsx"),
            )

            workbook_name = (
                safe_package_name(
                    (f"ReefVision_{package_type.title()}_Results_{run['run_name']}"),
                    limit=workbook_stem_limit,
                )
                + ".xlsx"
            )
            workbook = package / workbook_name
            _write_workbook(workbook, package_type, run, project, dataset, copied, warnings)
            copied.append((workbook, "workbook", "reefvision_generated"))
        _write_readme(package / "README.html", package_type, run, dataset, warnings)
        copied.append((package / "README.html", "report", "reefvision_generated"))
        manifest = _write_manifest(package, package_type, run, copied, warnings, timestamp)
        _write_checksums(package, manifest)
        verification = verify_results_package(package)
        if not verification["valid"]:
            raise RuntimeError(
                "Results package verification failed: " + "; ".join(verification["errors"])
            )
        package.rename(final_dir)
    archive = shutil.make_archive(
        str(final_dir), "zip", root_dir=final_dir.parent, base_dir=final_dir.name
    )
    return {
        "package_type": package_type,
        "package_name": final_dir.name,
        "archive_name": Path(archive).name,
        "archive_path": Path(archive),
        "archive_size_bytes": Path(archive).stat().st_size,
        "warnings": warnings,
    }


def verify_results_package(package: Path) -> dict[str, Any]:
    errors: list[str] = []
    manifest_path = package / "manifest.json"
    try:
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        return {"valid": False, "checked_files": 0, "errors": [f"Manifest unavailable: {exc}"]}
    if manifest.get("reefvision_results_package_schema") != PACKAGE_SCHEMA_VERSION:
        errors.append("Unsupported or missing package schema version")
    entries = manifest.get("files")
    if not isinstance(entries, list):
        return {"valid": False, "checked_files": 0, "errors": ["Manifest files must be a list"]}
    checked = 0
    expected: set[str] = set()
    for entry in entries:
        relative = str(entry.get("relative_path", ""))
        path = (package / relative).resolve()
        if (
            not relative
            or Path(relative).is_absolute()
            or not path.is_relative_to(package.resolve())
        ):
            errors.append(f"Unsafe manifest path: {relative}")
            continue
        expected.add(Path(relative).as_posix())
        if not path.is_file():
            errors.append(f"Missing file: {relative}")
            continue
        checked += 1
        if path.stat().st_size != entry.get("size_bytes"):
            errors.append(f"Size mismatch: {relative}")
        if sha256_file(path) != entry.get("sha256"):
            errors.append(f"Checksum mismatch: {relative}")
    actual = {
        item.relative_to(package).as_posix()
        for item in package.rglob("*")
        if item.is_file() and item.name not in {"manifest.json", "checksums.sha256"}
    }
    for relative in sorted(actual - expected):
        errors.append(f"Unlisted file: {relative}")
    return {"valid": not errors, "checked_files": checked, "errors": errors}


def _write_common_data(
    package: Path,
    run: Mapping[str, Any],
    project: Mapping[str, Any],
    dataset: Mapping[str, Any] | None,
    dataset_items: Sequence[Mapping[str, Any]],
    warnings: list[str],
) -> list[tuple[Path, str, str]]:
    data = package / "data"
    summary = {
        "schema_version": PACKAGE_SCHEMA_VERSION,
        "run_id": run["id"],
        "run_name": run["run_name"],
        "status": run["status"],
        "task": run["task"],
        "architecture": run["architecture"],
        "dataset_version_id": run.get("dataset_version_id"),
        "dataset_manifest_sha256": run.get("dataset_manifest_sha256"),
        "created_at": run.get("created_at"),
        "started_at": run.get("started_at"),
        "completed_at": run.get("completed_at"),
        "warnings": warnings,
    }
    provenance = _redact(
        {
            "schema_version": PACKAGE_SCHEMA_VERSION,
            "project": project,
            "dataset": dataset,
            "task": run["task"],
            "provider_id": run["provider_id"],
            "architecture": run["architecture"],
            "initial_checkpoint_source": run["initial_checkpoint_source"],
            "initial_checkpoint_model_id": run.get("initial_checkpoint_model_id"),
            "resolved_model_asset_id": run.get("resolved_model_asset_id"),
            "resolved_model_checksum": run.get("resolved_model_checksum"),
            "seed": run.get("random_seed"),
            "device": run.get("device"),
            "precision": run.get("precision"),
            "preflight_report": run.get("preflight_report", {}),
            "draft_snapshot": run.get("draft_snapshot", {}),
            "created_at": run.get("created_at"),
            "started_at": run.get("started_at"),
            "completed_at": run.get("completed_at"),
            "exported_at": datetime.now(UTC).isoformat(),
        }
    )
    effective_config = _redact(
        {
            "training": run.get("training_config", {}),
            "tuning": run.get("tuning_config", {}),
            "evaluation": run.get("evaluation_config", {}),
            "tiling": run.get("tiling_config", {}),
            "resources": run.get("resource_limits", {}),
        }
    )
    paths: list[tuple[Path, str, str]] = []
    for name, value in (
        ("run_summary.json", summary),
        ("provenance.json", provenance),
        ("effective_config.json", effective_config),
    ):
        path = data / name
        _write_json(path, value)
        paths.append((path, "data", "configuration"))
    yaml_path = data / "effective_config.yaml"
    yaml_path.write_text(yaml.safe_dump(effective_config, sort_keys=True), encoding="utf-8")
    paths.append((yaml_path, "data", "configuration"))
    if dataset:
        class_map = dataset.get("class_map", [])
        class_path = data / "class_map.csv"
        _write_csv(class_path, _rows(class_map), headers=None)
        paths.append((class_path, "data", "reefvision_generated"))
        split_path = data / "dataset_split_manifest.csv"
        _write_csv(split_path, dataset_items, headers=None)
        paths.append((split_path, "data", "reefvision_generated"))
    return paths


def _write_training_data(
    data: Path, run: Mapping[str, Any], warnings: list[str]
) -> list[tuple[Path, str, str]]:
    rows = _rows(run.get("native_results_rows", []))
    path = data / "epoch_metrics.csv"
    _write_csv(path, rows, headers=None)
    files = [(path, "data", "reefvision_generated")]
    if not rows:
        warnings.append("No persisted per-epoch metric rows were available")
    metrics = run.get("metrics", {})
    validation = data / "validation_metrics.csv"
    _write_csv(
        validation, [{"metric": key, "value": value} for key, value in _mapping(metrics).items()]
    )
    files.append((validation, "data", "reefvision_generated"))
    hyper = data / "hyperparameters.json"
    _write_json(hyper, run.get("training_config", {}))
    files.append((hyper, "data", "configuration"))
    hyper_yaml = data / "hyperparameters.yaml"
    hyper_yaml.write_text(
        yaml.safe_dump(run.get("training_config", {}), sort_keys=True), encoding="utf-8"
    )
    files.append((hyper_yaml, "data", "configuration"))
    per_class = _rows(run.get("per_class_metrics", []))
    if per_class:
        per_class_path = data / "per_class_metrics.csv"
        _write_csv(per_class_path, per_class)
        files.append((per_class_path, "data", "reefvision_generated"))
    for name, matrix in (
        ("confusion_matrix.csv", run.get("confusion_matrix")),
        ("confusion_matrix_normalized.csv", run.get("confusion_matrix_normalized")),
    ):
        matrix_rows = _rows(matrix)
        if matrix_rows:
            matrix_path = data / name
            _write_csv(matrix_path, matrix_rows)
            files.append((matrix_path, "data", "reefvision_generated"))
    return files


def _write_tuning_data(
    data: Path, run: Mapping[str, Any], warnings: list[str]
) -> list[tuple[Path, str, str]]:
    result = _mapping(run.get("tuning_results"))
    trials = _rows(result.get("trials", []))
    ranked = sorted(
        (trial for trial in trials if isinstance(trial.get("fitness"), (int, float))),
        key=lambda trial: (-float(trial["fitness"]), int(trial.get("trial_number", 0))),
    )
    all_rows = [_flatten_trial(trial) for trial in trials]
    parameter_rows = [
        {
            "trial_number": trial.get("trial_number"),
            "parameter": name,
            "value": value,
            "fitness": trial.get("fitness"),
            "status": trial.get("status"),
        }
        for trial in trials
        for name, value in _mapping(trial.get("hyperparameters")).items()
    ]
    epoch_rows = [
        {"trial_number": trial.get("trial_number"), **epoch}
        for trial in trials
        for epoch in _rows(trial.get("epoch_history", []))
    ]
    files: list[tuple[Path, str, str]] = []
    for name, rows in (
        ("all_trials.csv", all_rows),
        ("top_trials.csv", [_flatten_trial(item) for item in ranked[:10]]),
        ("parameter_fitness.csv", parameter_rows),
        ("trial_epoch_metrics.csv", epoch_rows),
    ):
        path = data / name
        _write_csv(path, rows)
        files.append((path, "data", "reefvision_generated"))
    for name, value in (
        ("tuning_summary.json", result),
        ("search_space.json", result.get("search_space", {})),
        ("best_trial.json", ranked[0] if ranked else None),
        ("best_hyperparameters.json", result.get("best_hyperparameters", {})),
    ):
        path = data / name
        _write_json(path, value)
        files.append((path, "data", "configuration"))
    best_yaml = data / "best_hyperparameters.yaml"
    best_yaml.write_text(
        yaml.safe_dump(result.get("best_hyperparameters", {}), sort_keys=True), encoding="utf-8"
    )
    files.append((best_yaml, "data", "configuration"))
    if not epoch_rows:
        warnings.append("The provider did not persist per-trial epoch histories")
    trial_artifacts = data / "trial_artifacts.csv"
    _write_csv(
        trial_artifacts,
        [
            {
                "trial_number": trial.get("trial_number"),
                **dict(artifact),
            }
            for trial in trials
            for artifact in _rows(trial.get("artifacts", []))
        ],
    )
    files.append((trial_artifacts, "data", "reefvision_generated"))
    return files


def _copy_artifacts(
    package: Path,
    artifacts: Sequence[tuple[Mapping[str, Any], Path]],
    options: Mapping[str, Any],
    warnings: list[str],
) -> list[tuple[Path, str, str]]:
    copied: list[tuple[Path, str, str]] = []
    seen: set[tuple[int, str]] = set()
    checkpoint_mode = str(options.get("checkpoint_mode", "best"))
    for index, (metadata, source) in enumerate(artifacts):
        category = str(metadata.get("category", "native"))
        name = source.name
        if category == "checkpoint":
            include = (
                checkpoint_mode == "retained"
                or (checkpoint_mode in {"best", "best_and_last"} and name.casefold() == "best.pt")
                or (checkpoint_mode == "best_and_last" and name.casefold() == "last.pt")
            )
            if not include:
                continue
            destination_dir = package / "checkpoints"
            source_type = "copied_checkpoint"
        elif category == "log":
            if not options.get("include_logs", True):
                continue
            destination_dir = package / "logs"
            source_type = "native"
        else:
            if not options.get("include_native_artifacts", True):
                continue
            destination_dir = (
                package / "plots" / "native"
                if category == "plot"
                else package / "artifacts" / "native"
            )
            source_type = "native"
        if not source.is_file() or source.is_symlink():
            warnings.append(f"Artifact unavailable or unsafe: {name}")
            continue
        identity = (source.stat().st_size, sha256_file(source))
        if identity in seen:
            continue
        seen.add(identity)
        destination_dir.mkdir(parents=True, exist_ok=True)
        destination_name_limit = safe_component_limit(
            destination_dir,
            requested=140,
        )

        destination = destination_dir / safe_package_name(
            name,
            limit=destination_name_limit,
        )
        if destination.exists():
            destination = destination.with_name(
                f"{destination.stem}_{index:04d}{destination.suffix}"
            )
        _stream_copy(source, destination)
        copied.append((destination, category, source_type))
    return copied


def _write_workbook(
    path: Path,
    package_type: str,
    run: Mapping[str, Any],
    project: Mapping[str, Any],
    dataset: Mapping[str, Any] | None,
    files: Sequence[tuple[Path, str, str]],
    warnings: Sequence[str],
) -> None:
    workbook = Workbook()
    active_sheet = workbook.active
    if active_sheet is not None:
        workbook.remove(active_sheet)
    workbook.properties.creator = "ReefVision"
    workbook.properties.title = f"ReefVision {package_type.title()} Results — {run['run_name']}"
    workbook.properties.subject = "Scientific computer-vision results"
    workbook.properties.description = "Recorded ReefVision result values and provenance"
    if package_type == "tuning":
        result = _mapping(run.get("tuning_results"))
        trials = _rows(result.get("trials", []))
        ranked = sorted(
            (trial for trial in trials if isinstance(trial.get("fitness"), (int, float))),
            key=lambda item: (-float(item["fitness"]), int(item.get("trial_number", 0))),
        )
        _key_value_sheet(
            workbook, "Summary", _tuning_summary(run, result, trials, ranked, warnings)
        )
        _table_sheet(workbook, "All Trials", [_flatten_trial(item) for item in trials])
        if ranked:
            _key_value_sheet(workbook, "Best Trial", _flatten_trial(ranked[0]))
            _table_sheet(workbook, "Top Trials", [_flatten_trial(item) for item in ranked[:10]])
        _search_space_sheet(workbook, result.get("search_space", {}))
        parameter_rows = [
            {
                "trial_number": trial.get("trial_number"),
                "parameter": key,
                "value": value,
                "fitness": trial.get("fitness"),
                "status": trial.get("status"),
            }
            for trial in trials
            for key, value in _mapping(trial.get("hyperparameters")).items()
        ]
        _table_sheet(workbook, "Parameter-Fitness", parameter_rows)
        epochs = [
            {"trial_number": trial.get("trial_number"), **row}
            for trial in trials
            for row in _rows(trial.get("epoch_history", []))
        ]
        _table_sheet(workbook, "Trial Epoch Metrics", epochs)
        durations = [
            float(item["duration_seconds"])
            for item in trials
            if isinstance(item.get("duration_seconds"), (int, float))
        ]
        if durations:
            _key_value_sheet(
                workbook,
                "Timing",
                {
                    "total_duration_seconds": sum(durations),
                    "average_trial_seconds": statistics.mean(durations),
                    "median_trial_seconds": statistics.median(durations),
                },
            )
    else:
        _key_value_sheet(workbook, "Run Summary", _training_summary(run, dataset, warnings))
        _table_sheet(workbook, "Epoch Metrics", _rows(run.get("native_results_rows", [])))
        _table_sheet(
            workbook,
            "Hyperparameters",
            [
                {"parameter": key, "effective_value": value}
                for key, value in _mapping(run.get("training_config")).items()
            ],
        )
        _table_sheet(
            workbook,
            "Validation Metrics",
            [
                {"metric": key, "value": value}
                for key, value in _mapping(run.get("metrics")).items()
            ],
        )
        _table_sheet(workbook, "Per-Class Metrics", _rows(run.get("per_class_metrics", [])))
        _table_sheet(workbook, "Confusion Matrix", _rows(run.get("confusion_matrix")))
        _table_sheet(
            workbook,
            "Confusion Matrix Norm",
            _rows(run.get("confusion_matrix_normalized")),
        )
        _key_value_sheet(
            workbook,
            "Hardware & Timing",
            {
                "device": run.get("device"),
                "precision": run.get("precision"),
                "started_at": run.get("started_at"),
                "completed_at": run.get("completed_at"),
                **_mapping(run.get("preflight_report")).get("resource_snapshot", {}),
            },
        )
    artifact_rows = [
        {
            "artifact_name": item.name,
            "category": category,
            "relative_package_path": item.relative_to(path.parent).as_posix(),
            "source_type": source,
            "file_size_bytes": item.stat().st_size,
            "sha256": sha256_file(item),
        }
        for item, category, source in files
        if item.is_file()
    ]
    _table_sheet(workbook, "Artifacts", artifact_rows)
    _key_value_sheet(
        workbook,
        "Provenance",
        {
            "project": project.get("name"),
            "project_id": project.get("id"),
            "dataset_version_id": run.get("dataset_version_id"),
            "dataset_checksum": run.get("dataset_manifest_sha256"),
            "task": run.get("task"),
            "provider": run.get("provider_id"),
            "architecture": run.get("architecture"),
            "model_checksum": run.get("resolved_model_checksum"),
            "seed": run.get("random_seed"),
            "created_at": run.get("created_at"),
            "exported_at": datetime.now(UTC).isoformat(),
        },
    )
    workbook.save(path)
    verified = load_workbook(path, read_only=True, data_only=False)
    if not verified.sheetnames:
        raise RuntimeError("Generated workbook contains no worksheets")
    verified.close()


def _key_value_sheet(workbook: Workbook, name: str, values: Mapping[str, Any]) -> None:
    sheet = workbook.create_sheet(name)
    sheet.append(["Field", "Value"])
    for key, value in values.items():
        sheet.append([_safe_cell(key), _safe_cell(_scalar(value))])
    _style_sheet(sheet, 2)


def _table_sheet(workbook: Workbook, name: str, rows: Sequence[Mapping[str, Any]]) -> None:
    if not rows:
        return
    headers = list(dict.fromkeys(key for row in rows for key in row))
    sheet = workbook.create_sheet(name[:31])
    sheet.append([_safe_cell(item) for item in headers])
    for row in rows:
        sheet.append([_safe_cell(_scalar(row.get(header))) for header in headers])
    _style_sheet(sheet, len(headers))


def _search_space_sheet(workbook: Workbook, value: Any) -> None:
    rows = []
    for name, definition in _mapping(value).items():
        fields = _mapping(definition)
        rows.append(
            {
                "parameter": name,
                "enabled": fields.get("enabled", True),
                "lower_bound": fields.get("min"),
                "upper_bound": fields.get("max"),
                "distribution": fields.get("distribution", "uniform"),
            }
        )
    _table_sheet(workbook, "Search Space", rows)


def _style_sheet(sheet: Any, column_count: int) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0F766E")
        cell.alignment = Alignment(wrap_text=True)
    for index in range(1, column_count + 1):
        values = [
            len(str(sheet.cell(row=row, column=index).value or ""))
            for row in range(1, min(sheet.max_row, 200) + 1)
        ]
        sheet.column_dimensions[get_column_letter(index)].width = min(
            max(values, default=8) + 2, 42
        )


def _write_manifest(
    package: Path,
    package_type: str,
    run: Mapping[str, Any],
    files: Sequence[tuple[Path, str, str]],
    warnings: Sequence[str],
    timestamp: datetime,
) -> dict[str, Any]:
    entries = []
    for path, category, source_type in files:
        if not path.is_file():
            continue
        relative = path.relative_to(package).as_posix()
        entries.append(
            {
                "relative_path": relative,
                "category": category,
                "source_type": source_type,
                "size_bytes": path.stat().st_size,
                "sha256": sha256_file(path),
                "media_type": mimetypes.guess_type(path.name)[0],
            }
        )
    manifest = {
        "reefvision_results_package_schema": PACKAGE_SCHEMA_VERSION,
        "package_type": package_type,
        "generated_at": timestamp.isoformat(),
        "reefvision_export_version": "1",
        "run_id": run["id"],
        "reproducibility_id": run["id"],
        "warnings": list(dict.fromkeys(warnings)),
        "files": entries,
        "excluded_from_manifest": ["manifest.json", "checksums.sha256"],
    }
    _write_json(package / "manifest.json", manifest)
    return manifest


def _write_checksums(package: Path, manifest: Mapping[str, Any]) -> None:
    rows = [f"{entry['sha256']}  {entry['relative_path']}" for entry in manifest["files"]]
    rows.append(f"{sha256_file(package / 'manifest.json')}  manifest.json")
    (package / "checksums.sha256").write_text("\n".join(rows) + "\n", encoding="utf-8")


def _write_standard_plots(package: Path, run: Mapping[str, Any]) -> list[tuple[Path, str, str]]:
    output = package / "plots" / "reefvision"
    series: list[tuple[str, list[float]]] = []
    tuning = _mapping(run.get("tuning_results"))
    if tuning:
        trials = _rows(tuning.get("trials", []))
        fitness = [
            float(item["fitness"])
            for item in trials
            if isinstance(item.get("fitness"), (int, float))
        ]
        if fitness:
            series.append(("fitness_vs_trial", fitness))
            current = float("-inf")
            best = []
            for value in fitness:
                current = max(current, value)
                best.append(current)
            series.append(("best_fitness_so_far", best))
    else:
        for name, values in _mapping(run.get("metrics_history")).items():
            numeric = (
                [float(value) for value in values if isinstance(value, (int, float))]
                if isinstance(values, list)
                else []
            )
            if numeric and name.casefold() not in {"epoch", "time"}:
                series.append((safe_package_name(name.replace("/", "_"), limit=70), numeric))
    if not series:
        return []
    output.mkdir(parents=True, exist_ok=True)
    files = []
    for name, values in series:
        low, high = min(values), max(values)
        spread = high - low or 1.0
        coordinates = []
        for index, value in enumerate(values):
            x_value = 50 + index * 700 / max(len(values) - 1, 1)
            y_value = 350 - (value - low) * 280 / spread
            coordinates.append(f"{x_value:.3f},{y_value:.3f}")
        points = " ".join(coordinates)
        path = output / f"{name}.svg"
        title = escape(name.replace("_", " ").title())
        svg = (
            "<svg xmlns='http://www.w3.org/2000/svg' width='800' height='420' "
            f"viewBox='0 0 800 420' role='img' aria-label='{escape(name)}'>"
            "<rect width='800' height='420' fill='white'/>"
            "<path d='M50 35V350H765' stroke='#64748b' fill='none'/>"
            f"<text x='50' y='24' font-family='sans-serif' font-size='18'>{title}</text>"
            f"<polyline points='{points}' fill='none' stroke='#0f766e' stroke-width='3'/>"
            "<text x='50' y='385' font-family='sans-serif' font-size='12'>"
            "Recorded values; no smoothing</text></svg>"
        )
        path.write_text(
            svg,
            encoding="utf-8",
        )
        files.append((path, "plot", "reefvision_generated"))
    return files


def _write_readme(
    path: Path,
    package_type: str,
    run: Mapping[str, Any],
    dataset: Mapping[str, Any] | None,
    warnings: Sequence[str],
) -> None:
    def escaped(value: Any) -> str:
        return str(value).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    items = (
        "".join(f"<li>{escaped(item)}</li>" for item in dict.fromkeys(warnings))
        or "<li>None recorded</li>"
    )
    dataset_name = escaped((dataset or {}).get("name", run.get("dataset_version_id")))
    html = (
        "<!doctype html><html><meta charset='utf-8'>"
        "<title>ReefVision Results Package</title><body>"
        f"<h1>ReefVision {package_type.title()} Results Package</h1><dl>"
        f"<dt>Run</dt><dd>{escaped(run['run_name'])}</dd>"
        f"<dt>Run ID</dt><dd>{escaped(run['id'])}</dd>"
        f"<dt>Status</dt><dd>{escaped(run['status'])}</dd>"
        f"<dt>Task</dt><dd>{escaped(run['task'])}</dd>"
        f"<dt>Dataset</dt><dd>{dataset_name}</dd></dl>"
        "<h2>Contents</h2><p>The workbook summarizes recorded values. "
        "The data directory contains machine-readable tables and provenance. "
        "Native provider artifacts remain unmodified copies. manifest.json and "
        "checksums.sha256 support verification.</p>"
        f"<h2>Warnings</h2><ul>{items}</ul></body></html>"
    )
    path.write_text(html, encoding="utf-8")


def _training_summary(
    run: Mapping[str, Any], dataset: Mapping[str, Any] | None, warnings: Sequence[str]
) -> dict[str, Any]:
    rows = _rows(run.get("native_results_rows", []))
    return {
        "run_name": run["run_name"],
        "run_id": run["id"],
        "status": run["status"],
        "task": run["task"],
        "architecture": run["architecture"],
        "dataset_version": (dataset or {}).get("name"),
        "dataset_checksum": run.get("dataset_manifest_sha256"),
        "configured_epochs": _mapping(run.get("training_config")).get("epochs"),
        "completed_epochs": len(rows),
        "final_epoch": rows[-1].get("epoch") if rows else None,
        "best_checkpoint": "best.pt" if run.get("result_checkpoint_path") else None,
        "started_at": run.get("started_at"),
        "completed_at": run.get("completed_at"),
        "warnings": "; ".join(warnings),
    }


def _tuning_summary(
    run: Mapping[str, Any],
    result: Mapping[str, Any],
    trials: Sequence[Mapping[str, Any]],
    ranked: Sequence[Mapping[str, Any]],
    warnings: Sequence[str],
) -> dict[str, Any]:
    statuses = [str(item.get("status")) for item in trials]
    return {
        "run_name": run["run_name"],
        "run_id": run["id"],
        "status": run["status"],
        "requested_max_trials": result.get("max_trials"),
        "completed_trials": statuses.count("completed"),
        "failed_trials": statuses.count("failed"),
        "epochs_per_trial": result.get("epochs_per_trial"),
        "patience": result.get("patience"),
        "best_trial": ranked[0].get("trial_number") if ranked else None,
        "best_fitness": ranked[0].get("fitness") if ranked else None,
        "fitness_definition": result.get("fitness_definition"),
        "warnings": "; ".join(warnings),
    }


def _flatten_trial(trial: Mapping[str, Any]) -> dict[str, Any]:
    return {
        "trial_number": trial.get("trial_number"),
        "trial_id": trial.get("trial_id", trial.get("trial_number")),
        "status": trial.get("status"),
        "fitness": trial.get("fitness"),
        "duration_seconds": trial.get("duration_seconds"),
        "epoch_count": len(_rows(trial.get("epoch_history", []))),
        **{f"metric__{key}": value for key, value in _mapping(trial.get("metrics")).items()},
        **{
            f"parameter__{key}": value
            for key, value in _mapping(trial.get("hyperparameters")).items()
        },
    }


def _write_csv(
    path: Path, rows: Iterable[Mapping[str, Any]], headers: Sequence[str] | None = None
) -> None:
    materialized = list(rows)
    columns = list(headers or dict.fromkeys(key for row in materialized for key in row))
    with path.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=columns, extrasaction="ignore")
        writer.writeheader()
        for row in materialized:
            writer.writerow({key: _csv_cell(row.get(key)) for key in columns})


def _csv_cell(value: Any) -> Any:
    value = _scalar(value)
    return _safe_cell(value)


def _safe_cell(value: Any) -> Any:
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _scalar(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, float, bool, datetime)):
        return value
    return json.dumps(value, ensure_ascii=False, sort_keys=True, allow_nan=False)


def _redact(value: Any) -> Any:
    if isinstance(value, Mapping):
        return {
            str(key): "[redacted]" if SECRET_KEYS.search(str(key)) else _redact(item)
            for key, item in value.items()
        }
    if isinstance(value, list):
        return [_redact(item) for item in value]
    return value


def _write_json(path: Path, value: Any) -> None:
    path.write_text(
        json.dumps(value, ensure_ascii=False, indent=2, sort_keys=True, allow_nan=False) + "\n",
        encoding="utf-8",
    )


def _stream_copy(source: Path, destination: Path) -> None:
    with source.open("rb") as source_stream, destination.open("xb") as destination_stream:
        shutil.copyfileobj(source_stream, destination_stream, length=CHUNK_SIZE)


def _rows(value: Any) -> list[dict[str, Any]]:
    return (
        [dict(item) for item in value if isinstance(item, Mapping)]
        if isinstance(value, list)
        else []
    )


def _mapping(value: Any) -> dict[str, Any]:
    return dict(value) if isinstance(value, Mapping) else {}
